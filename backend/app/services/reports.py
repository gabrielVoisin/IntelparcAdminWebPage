"""Report generation (PDF with ReportLab, Excel with openpyxl).

Each report is one SQL query returning rows for a site (or every site) and a date range.
Column titles are translated; values stay raw (cents converted to currency units).
"""

from __future__ import annotations

import io
import uuid
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta

from sqlalchemy import text
from sqlalchemy.orm import Session

from app.i18n import translate


@dataclass(frozen=True)
class ReportDef:
    sql: str
    columns: tuple[str, ...]          # translation keys "report.col.<name>"
    money_columns: tuple[str, ...] = ()
    plate_columns: tuple[str, ...] = ()


SCOPE = "(CAST(:site_id AS uuid) IS NULL OR {col} = CAST(:site_id AS uuid))"

REPORTS: dict[str, ReportDef] = {
    "revenue": ReportDef(
        f"""SELECT date(paid_at AT TIME ZONE 'Europe/Paris') AS day,
                   CASE WHEN subscription_id IS NOT NULL THEN 'subscription'
                        WHEN reservation_id IS NOT NULL THEN 'reservation' ELSE 'hourly_daily' END AS source,
                   count(*) AS payments, sum(amount_cents) AS amount,
                   sum(round(amount_cents * vat_rate_bp / (10000.0 + vat_rate_bp)))::int AS vat
            FROM payments WHERE status = 'succeeded' AND paid_at >= :from_ts AND paid_at < :to_ts AND {SCOPE.format(col='site_id')}
            GROUP BY 1, 2 ORDER BY 1, 2""",
        ("day", "source", "payments", "amount", "vat"), money_columns=("amount", "vat")),
    "transactions": ReportDef(
        f"""SELECT p.paid_at, p.amount_cents AS amount, p.method, p.status, ps.plate_display AS plate, p.provider_ref
            FROM payments p LEFT JOIN parking_sessions ps ON ps.id = p.session_id
            WHERE p.created_at >= :from_ts AND p.created_at < :to_ts AND {SCOPE.format(col='p.site_id')}
            ORDER BY p.created_at""",
        ("paid_at", "amount", "method", "status", "plate", "provider_ref"), money_columns=("amount",), plate_columns=("plate",)),
    "occupancy": ReportDef(
        f"""SELECT s.name AS site, z.name AS zone, count(b.id) AS bays,
                   count(b.id) FILTER (WHERE b.is_occupied) AS occupied_now,
                   (SELECT count(*) FROM parking_sessions p JOIN bays b2 ON b2.id = p.bay_id
                     WHERE b2.zone_id = z.id AND p.entered_at >= :from_ts AND p.entered_at < :to_ts) AS stays,
                   round(100.0 * (SELECT coalesce(sum(extract(epoch FROM least(coalesce(p.exited_at, now()), :to_ts) - greatest(p.entered_at, :from_ts))), 0)
                     FROM parking_sessions p JOIN bays b2 ON b2.id = p.bay_id WHERE b2.zone_id = z.id
                       AND p.entered_at < :to_ts AND coalesce(p.exited_at, now()) > :from_ts)
                     / nullif(count(b.id) * extract(epoch FROM (:to_ts - :from_ts)), 0), 1) AS avg_occupancy_pct
            FROM zones z JOIN sites s ON s.id = z.site_id JOIN bays b ON b.zone_id = z.id
            WHERE {SCOPE.format(col='s.id')} GROUP BY s.name, z.id, z.name, z.sort_order ORDER BY s.name, z.sort_order""",
        ("site", "zone", "bays", "occupied_now", "stays", "avg_occupancy_pct")),
    "entries_exits": ReportDef(
        f"""SELECT p.entered_at, p.exited_at, p.plate_display AS plate, z.name AS zone, b.number AS bay,
                   CASE WHEN p.subscription_id IS NOT NULL THEN 'subscriber' WHEN p.reservation_id IS NOT NULL THEN 'reservation' ELSE 'occasional' END AS customer_type
            FROM parking_sessions p LEFT JOIN bays b ON b.id = p.bay_id LEFT JOIN zones z ON z.id = b.zone_id
            WHERE p.entered_at >= :from_ts AND p.entered_at < :to_ts AND {SCOPE.format(col='p.site_id')} ORDER BY p.entered_at""",
        ("entered_at", "exited_at", "plate", "zone", "bay", "customer_type"), plate_columns=("plate",)),
    "parking_duration": ReportDef(
        f"""WITH d AS (SELECT extract(epoch FROM coalesce(exited_at, now()) - entered_at) / 60 AS minutes,
                          CASE WHEN subscription_id IS NOT NULL THEN 'subscriber' WHEN reservation_id IS NOT NULL THEN 'reservation' ELSE 'occasional' END AS profile
                   FROM parking_sessions WHERE entered_at >= :from_ts AND entered_at < :to_ts AND {SCOPE.format(col='site_id')})
            SELECT profile, count(*) AS stays, round(avg(minutes)) AS avg_minutes,
                   round(percentile_cont(0.5) WITHIN GROUP (ORDER BY minutes)) AS median_minutes,
                   count(*) FILTER (WHERE minutes < 60) AS under_1h, count(*) FILTER (WHERE minutes >= 60 AND minutes < 240) AS from_1h_to_4h,
                   count(*) FILTER (WHERE minutes >= 240) AS over_4h
            FROM d GROUP BY profile ORDER BY profile""",
        ("profile", "stays", "avg_minutes", "median_minutes", "under_1h", "from_1h_to_4h", "over_4h")),
    "peak_hours": ReportDef(
        f"""SELECT extract(isodow FROM entered_at AT TIME ZONE 'Europe/Paris')::int AS weekday,
                   extract(hour FROM entered_at AT TIME ZONE 'Europe/Paris')::int AS hour, count(*) AS entries
            FROM parking_sessions WHERE entered_at >= :from_ts AND entered_at < :to_ts AND {SCOPE.format(col='site_id')}
            GROUP BY 1, 2 ORDER BY entries DESC, 1, 2""",
        ("weekday", "hour", "entries")),
    "subscriptions": ReportDef(
        f"""SELECT c.reference AS customer, p.code AS plan, s.price_cents AS amount, s.starts_on, s.ends_on, s.status,
                   s.auto_renew, s.payment_overdue
            FROM subscriptions s JOIN customers c ON c.id = s.customer_id JOIN subscription_plans p ON p.id = s.plan_id
            WHERE s.starts_on <= :to_date AND s.ends_on >= :from_date AND {SCOPE.format(col='s.site_id')} ORDER BY s.ends_on""",
        ("customer", "plan", "amount", "starts_on", "ends_on", "status", "auto_renew", "payment_overdue"), money_columns=("amount",)),
    "payments": ReportDef(
        f"""SELECT method, status, count(*) AS payments, sum(amount_cents) AS amount
            FROM payments WHERE created_at >= :from_ts AND created_at < :to_ts AND {SCOPE.format(col='site_id')}
            GROUP BY 1, 2 ORDER BY 1, 2""",
        ("method", "status", "payments", "amount"), money_columns=("amount",)),
    "anomalies": ReportDef(
        f"""SELECT a.detected_at, a.reason, a.plate_display AS plate, z.name AS zone, b.number AS bay, a.status,
                   f.amount_cents AS fine_amount, f.status AS fine_status
            FROM payment_alarms a LEFT JOIN bays b ON b.id = a.bay_id LEFT JOIN zones z ON z.id = b.zone_id
            LEFT JOIN fines f ON f.alarm_id = a.id
            WHERE a.detected_at >= :from_ts AND a.detected_at < :to_ts AND {SCOPE.format(col='a.site_id')} ORDER BY a.detected_at""",
        ("detected_at", "reason", "plate", "zone", "bay", "status", "fine_amount", "fine_status"),
        money_columns=("fine_amount",), plate_columns=("plate",)),
    "equipment": ReportDef(
        f"""SELECT c.code AS controller, e.code AS device, e.priority, e.message_key AS message, e.raised_at, e.cleared_at,
                   e.acknowledged_at, round(extract(epoch FROM coalesce(e.cleared_at, now()) - e.raised_at) / 60) AS downtime_minutes
            FROM device_events e JOIN controllers c ON c.id = e.controller_id
            WHERE e.raised_at >= :from_ts AND e.raised_at < :to_ts AND {SCOPE.format(col='c.site_id')} ORDER BY e.raised_at""",
        ("controller", "device", "priority", "message", "raised_at", "cleared_at", "acknowledged_at", "downtime_minutes")),
}


def _mask_plate(plate: str | None) -> str | None:
    if not plate:
        return plate
    return plate[:3] + "*" * max(0, len(plate) - 3)


def fetch_rows(db: Session, report_type: str, site_id: uuid.UUID | None, start: date, end: date, anonymise: bool) -> tuple[ReportDef, list[dict]]:
    definition = REPORTS[report_type]
    params = {
        "site_id": site_id,
        "from_ts": datetime.combine(start, time.min).astimezone(),
        "to_ts": datetime.combine(end + timedelta(days=1), time.min).astimezone(),
        "from_date": start,
        "to_date": end,
    }
    rows = [dict(r._mapping) for r in db.execute(text(definition.sql), params)]
    for row in rows:
        for col in definition.money_columns:
            if row.get(col) is not None:
                row[col] = round(row[col] / 100, 2)
        if anonymise:
            for col in definition.plate_columns:
                row[col] = _mask_plate(row.get(col))
    return definition, rows


def _cell(value: object) -> object:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    return value


def render_xlsx(title: str, headers: list[str], keys: tuple[str, ...], rows: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(headers)
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="16181D")
    for row in rows:
        ws.append([_cell(row.get(k)) for k in keys])
    for i, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(header) + 4)
    ws.freeze_panes = "A4"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def render_pdf(title: str, subtitle: str, headers: list[str], keys: tuple[str, ...], rows: list[dict]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=15 * mm)
    styles = getSampleStyleSheet()

    def fmt(v: object) -> str:
        if v is None:
            return ""
        if isinstance(v, datetime):
            return v.strftime("%d.%m.%Y %H:%M")
        if isinstance(v, date):
            return v.strftime("%d.%m.%Y")
        return str(v)

    data = [headers] + [[fmt(r.get(k)) for k in keys] for r in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16181D")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F7F3")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9CCC4")),
        ])
    )
    doc.build([Paragraph(title, styles["Title"]), Paragraph(subtitle, styles["Normal"]), Spacer(1, 6 * mm), table])
    return buffer.getvalue()


def build(db: Session, report_type: str, fmt: str, site_id: uuid.UUID | None, site_name: str,
          start: date, end: date, locale: str, anonymise: bool) -> tuple[bytes, str]:
    definition, rows = fetch_rows(db, report_type, site_id, start, end, anonymise)
    title = translate(f"report.{report_type}.title", locale)
    headers = [translate(f"report.col.{c}", locale) for c in definition.columns]
    subtitle = f"{site_name} · {start:%d.%m.%Y} → {end:%d.%m.%Y}"
    if fmt == "xlsx":
        return render_xlsx(title, headers, definition.columns, rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return render_pdf(title, subtitle, headers, definition.columns, rows), "application/pdf"
